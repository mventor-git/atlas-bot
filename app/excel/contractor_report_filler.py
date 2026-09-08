"""
Contractor Report Filler — fills the contractor_period_report template
with aggregated data for a specific contractor over a date range.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

if TYPE_CHECKING:
    from app.models.database import Report, ReportItem

logger = logging.getLogger(__name__)


class ContractorReportFiller:
    """Fills the contractor period report template with data."""

    # ── Styles (matching the template) ──
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    alt_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def __init__(self, template_path: str | Path) -> None:
        self._template_path = Path(template_path)

    def fill(
        self,
        contractor_name: str,
        start_date: str,
        end_date: str,
        entries: list[dict[str, Any]],
        output_path: str | Path,
    ) -> Path:
        """Fill the template and save to ``output_path``.

        Args:
            contractor_name: The contractor's display name.
            start_date: Period start (YYYY-MM-DD).
            end_date: Period end (YYYY-MM-DD).
            entries: List of dicts with keys:
                date, workers, zone, details, added_by, role.
            output_path: Where to save the filled workbook.

        Returns:
            Path to the saved workbook.
        """
        wb = load_workbook(self._template_path)
        ws = wb.active

        # ── Row 2: Info line ──
        ws["A2"] = f"Contractor: {contractor_name}  |  Period: {start_date} to {end_date}"

        # ── Column letters for data (B=Date, C=Workers, D=Zone, E=Details, F=Added By, G=Role) ──
        data_cols = ["B", "C", "D", "E", "F", "G"]

        # Remove sample rows (row 5 and 6) — we'll write real data
        for row_num in range(5, 7):
            for col in data_cols:
                ws[f"{col}{row_num}"] = None
                cell = ws[f"{col}{row_num}"]
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment()
                cell.fill = PatternFill()
                cell.border = Border()

        # Clear summary row
        ws.merge_cells("A7:G7")
        ws["A7"] = None
        ws["A7"].font = Font(name="Calibri", size=11, bold=True)
        ws["A7"].alignment = Alignment()

        # ── Write data rows starting at row 5 ──
        total_workers = 0
        for i, entry in enumerate(entries):
            row_num = 5 + i
            w = entry.get("workers", 0) or 0
            total_workers += w
            is_alt = i % 2 == 1  # alternating rows

            values = [
                entry.get("date", ""),
                str(w),
                entry.get("zone", ""),
                entry.get("details", ""),
                entry.get("added_by", ""),
                entry.get("role", ""),
            ]
            for col, val in zip(data_cols, values):
                cell = ws[f"{col}{row_num}"]
                cell.value = val
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = self.left_wrap if col in ("D", "E", "F") else self.center
                cell.border = self.thin_border
                if is_alt:
                    cell.fill = self.alt_fill

            # Ensure row height for wrapped text
            ws.row_dimensions[row_num].height = 20

        # ── Summary row ──
        summary_row = 5 + len(entries)
        ws.merge_cells(f"A{summary_row}:G{summary_row}")
        ws[f"A{summary_row}"] = f"Summary: {len(entries)} entries, {total_workers} total workers"
        ws[f"A{summary_row}"].font = Font(name="Calibri", size=11, bold=True)
        ws[f"A{summary_row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"A{summary_row}"].border = self.thin_border

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output))
        logger.info("Contractor report saved to %s", output)
        return output
